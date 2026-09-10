from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import date

doc = Document()

for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

title = doc.add_heading("Stay Sticky Privacy Policy", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = meta.add_run(
    f"Effective date: {date.today().isoformat()}\nLast updated: {date.today().isoformat()}"
)
run.italic = True

doc.add_paragraph(
    'Stay Sticky ("Stay Sticky," "we," "us," or "our") is a Chrome extension and companion web library '
    "that lets you pin sticky notes on webpages and optionally sync them to your account. "
    "This Privacy Policy explains what information we collect, how we use it, and the choices you have."
)

doc.add_heading("1. Who this policy covers", level=1)
doc.add_paragraph(
    "This policy applies to the Stay Sticky Chrome extension, the companion website "
    "(including staysticky-web.vercel.app and any related domains we operate), and related backend services."
)

doc.add_heading("2. Information we collect", level=1)

doc.add_heading("2.1 Information you provide", level=2)
for i in [
    "Account information when you sign in with Google (for example, your name, email address, and profile photo URL) through Firebase Authentication.",
    "Sticky note content you create (note text, colors, size/position, minimize state, tags, project assignments, and related metadata you choose to save).",
    "Optional page context associated with a note (page URL, page title, page path key, and—if you enable snapshots—surrounding text or a page snapshot you allow us to store).",
    "Account preferences such as sync, auto-grouping, and snapshot settings.",
]:
    doc.add_paragraph(i, style="List Bullet")

doc.add_heading("2.2 Information stored on your device", level=2)
doc.add_paragraph(
    "By default, notes are stored locally in your browser using Chrome's extension storage "
    "(chrome.storage.local). Local notes remain on your device until you delete them or clear extension data."
)

doc.add_heading("2.3 Information synced to our servers (when you connect an account)", level=2)
doc.add_paragraph(
    "If you sign in and enable sync, we store your notes, projects, and settings in Google Firebase "
    "(Cloud Firestore) under your user account so you can access them from the companion website and across devices."
)

doc.add_heading("2.4 Technical / usage information", level=2)
doc.add_paragraph(
    "Our hosting and authentication providers (for example Vercel and Google Firebase) may automatically "
    "process standard technical data such as IP address, browser type, device information, timestamps, "
    "and security logs needed to operate and protect the service. We do not use this to build advertising profiles."
)

doc.add_heading("3. How we use information", level=1)
for u in [
    "Provide, maintain, and improve the extension and web library (display notes, sync, search, project grouping, and draft extractive summaries generated from your note text without selling your content to advertisers).",
    "Authenticate you and keep your data associated with your account.",
    "Respond to support requests and communicate about the service when needed.",
    "Protect security, prevent abuse, and comply with legal obligations.",
]:
    doc.add_paragraph(u, style="List Bullet")

doc.add_heading("4. How we share information", level=1)
doc.add_paragraph(
    "We do not sell your personal information. We do not share your note content with third parties for advertising."
)
doc.add_paragraph(
    "We use service providers who process data on our behalf to run Stay Sticky:"
)
for p in [
    "Google Firebase (Authentication and Cloud Firestore) — account and synced note/project storage.",
    "Vercel — hosting the companion website.",
    "Google (as your sign-in provider) when you choose Google Sign-In.",
]:
    doc.add_paragraph(p, style="List Bullet")
doc.add_paragraph(
    "We may disclose information if required by law, legal process, or to protect the rights, safety, "
    "and security of users or the public."
)

doc.add_heading("5. Chrome extension permissions", level=1)
doc.add_paragraph(
    "Stay Sticky requests permissions needed for its core features, including storing notes, interacting with tabs, "
    "and injecting the note UI on pages you visit. Host access is used to show notes on webpages. "
    "We do not use permissions to scrape unrelated browsing history for advertising."
)

doc.add_heading("6. Data retention", level=1)
doc.add_paragraph(
    "Local notes remain until you delete them or remove/clear the extension. "
    "Synced account data remains until you delete notes/projects, disconnect sync and delete data, "
    "or request account deletion. Backup or log copies held by providers may persist for a limited period "
    "consistent with their policies and our operational needs."
)

doc.add_heading("7. Your choices and rights", level=1)
for r in [
    "Use Stay Sticky without signing in (notes stay local only).",
    "Turn sync and related settings on or off in Account & sync.",
    "Edit or delete individual notes from the extension or web library.",
    "Sign out of your account on the website.",
    "Uninstall the extension and/or clear extension storage in Chrome.",
    "Request access to or deletion of your account data by contacting us (see Contact).",
]:
    doc.add_paragraph(r, style="List Bullet")
doc.add_paragraph(
    "Depending on where you live (for example under GDPR or CCPA/CPRA), you may have additional rights "
    "regarding access, correction, deletion, portability, or opting out of certain processing. "
    "We will respond to verifiable requests as required by applicable law."
)

doc.add_heading("8. Children's privacy", level=1)
doc.add_paragraph(
    "Stay Sticky is not directed to children under 13 (or the minimum age required in your jurisdiction), "
    "and we do not knowingly collect personal information from children."
)

doc.add_heading("9. International transfers", level=1)
doc.add_paragraph(
    "Our providers may process data in the United States or other countries. "
    "By using Stay Sticky, you understand that your information may be transferred to and processed in "
    "locations that may have different data-protection laws than your home country."
)

doc.add_heading("10. Security", level=1)
doc.add_paragraph(
    "We use industry-standard safeguards provided by our platforms (including authenticated access controls "
    "and Firebase security rules that restrict each signed-in user to their own data). "
    "No method of transmission or storage is 100% secure."
)

doc.add_heading("11. Third-party websites", level=1)
doc.add_paragraph(
    "Notes can be attached to third-party websites you visit. Those sites have their own privacy practices. "
    "Stay Sticky does not control third-party site policies."
)

doc.add_heading("12. Changes to this policy", level=1)
doc.add_paragraph(
    'We may update this Privacy Policy from time to time. We will post the updated version with a revised '
    '"Last updated" date. Continued use of Stay Sticky after changes means you accept the updated policy.'
)

doc.add_heading("13. Contact", level=1)
doc.add_paragraph(
    "If you have questions about this Privacy Policy or want to request access or deletion of your data, contact:"
)
doc.add_paragraph("Email: pulugurtha.inesh@gmail.com")
doc.add_paragraph("Project: Stay Sticky (Chrome extension + companion web library)")
doc.add_paragraph("Website: https://staysticky-web.vercel.app")
doc.add_paragraph("Source / support: https://github.com/ineshpul/staysticky")

footer = doc.add_paragraph()
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run("End of Privacy Policy")
fr.italic = True

out = r"C:\Users\pulug\staysticky\Stay-Sticky-Privacy-Policy.docx"
doc.save(out)
print(out)
